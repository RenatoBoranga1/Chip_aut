"""V16 adapter and existing publication workflow, using generated isolated fixtures."""

import io
import sqlite3
from copy import deepcopy
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import Workbook
from test_matching import motorcycle
from test_review_queue import ad, coverage, import_base
from test_scanner_versions import publish

from database.review_repository import ReviewRepository
from scanner_base.adapters import (
    ApplicationGeneralScannerBaseAdapter,
    detect_adapter,
    parse_scanner_workbook,
)
from scanner_base.application_diff import application_differences
from scanner_base.excel_reader import SheetData, WorkbookData, read_excel
from services.dashboard_service import DashboardConfig, DashboardService
from services.import_service import load_rules
from services.management_metrics_service import ZONE, ManagementMetricsService
from services.scanner_service import ScannerService
from services.scanner_validation import validate_upload

HEADERS = [
    "VERS.",
    "MONTADORA",
    "MODELO",
    "ANO",
    "SISTEMA",
    "CABO",
    "TIPO DE TESTE",
    "NOVO SISTEMA",
    "VÍDEO",
    "TABELA FIPE",
    "IMOBILIZADOR",
    "FUNÇÕES AVANÇADAS",
    "SCOOTER",
]


def application(**changes):
    return {
        **dict(
            zip(
                HEADERS,
                ["V1.0", "BMW", "F 900 R", 2025, "ABS", "MX-C1", "EM CAMPO", "", "SIM", "NÃO", "NÃO", "", "NÃO"],
            )
        ),
        **changes,
    }


def content(rows=None, headers=None, extra=False, title="APLICACAO GERAL"):
    wb = Workbook()
    ws = wb.active
    ws.title = title
    headers = headers or HEADERS
    ws.append(headers)
    for r in rows or [application()]:
        ws.append([r.get(h) for h in headers])
    if extra:
        ws = wb.create_sheet("APLICAÇÃO V16")
        ws.append(HEADERS[:6])
        ws.append(["V15.0", "BMW", "NOT IN CONSOLIDATED", 2026, "PAINEL", "C2"])
        ws = wb.create_sheet("VERSOES")
        ws.append([None, "V14.0"])
        ws.append(["APLICAÇÃO", 1])
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()


def parse(rows=None, **kwargs):
    return validate_upload(content(rows, **kwargs), "new.xlsx")


def test_adapter_detects_and_preserves_legacy():
    from test_scanner import HEADERS as OLD_HEADERS
    from test_scanner import row

    book = WorkbookData([SheetData("Base", rows=[(1, OLD_HEADERS), (2, row())])], False)
    assert detect_adapter(book).format_name == "LEGACY"
    assert parse_scanner_workbook(book, load_rules()).records[0].release == "Sim"


def test_consolidated_source_and_metadata_not_double_counted():
    base, _, report, _ = parse([application(), application(SISTEMA="INJEÇÃO")], extra=True)
    assert report["format"] == "APPLICATION_GENERAL"
    assert report["applications"] == 2 and report["unique_vehicles"] == 1
    assert all(m.model == "F 900 R" for m in base.motorcycles)
    assert report["metadata"]["complementary_releases"][0]["name_version_mismatch"]
    assert report["issues"]["SUMMARY_TOTAL_DIFFERS"] == 1


@pytest.mark.parametrize("missing", HEADERS)
def test_missing_header_rejected(missing):
    with pytest.raises(ValueError):
        parse(headers=[h for h in HEADERS if h != missing])


def test_reordered_columns_and_duplicate_headers():
    b, _, _, _ = parse(headers=list(reversed(HEADERS)))
    assert b.records[0].key == "BMW|F900R|2025"
    with pytest.raises(ValueError, match="duplicado"):
        parse(headers=HEADERS + ["MONTADORA"])


def test_no_primary_does_not_import_release_as_full_base():
    with pytest.raises(ValueError, match="consolidada"):
        parse(title="APLICAÇÃO V16")


def test_multiple_systems_years_and_cables():
    b, _, r, _ = parse(
        [application(), application(SISTEMA="INJEÇÃO"), application(CABO="MX-C2"), application(ANO=2026)]
    )
    assert len(b.motorcycles) == 2 and len(b.records) == 4
    assert r["application_keys"] == 4 and r["unique_systems"] == 2
    assert {r.cable for r in b.records} == {"MX-C1", "MX-C2"}


def test_exact_duplicate_and_attribute_variant_preserved():
    b, _, r, _ = parse([application(), application(), application(**{"VÍDEO": "NÃO"})])
    assert len(b.records) == 3 and r["unique_applications"] == 2 and r["application_keys"] == 1
    assert b.records[1].duplicate_of == 0
    assert r["issues"]["APPLICATION_VARIANT"] == 1
    assert len(b.motorcycles) == 1


@pytest.mark.parametrize("year", [None, "2024/2025", "abc", 1800, 2200, 2025.5])
def test_invalid_year_reported_not_corrected(year):
    b, _, r, _ = parse([application(), application(ANO=year)])
    assert b.rejected_rows == 1 and r["invalid_records"] == 1 and not r["publishable"]


def test_configurable_year_range():
    book = read_excel(io.BytesIO(content([application(ANO=2024), application(ANO=2025)])))
    base = ApplicationGeneralScannerBaseAdapter({"min_year": 2025, "max_year": 2026}).parse(book, load_rules())
    assert base.rejected_rows == 1 and base.records[0].year == 2025


def test_optional_blank_and_unknown_flags_preserve_raw():
    b, _, r, _ = parse(
        [application(**{"TIPO DE TESTE": "", "TABELA FIPE": "Verificar nome", "FUNÇÕES AVANÇADAS": None})]
    )
    record = b.records[0]
    assert record.application_attributes["fipe"] is None
    assert record.raw_fields["fipe"] == "Verificar nome"
    assert record.test_type is None and record.application_attributes["advanced_functions"] is None
    assert r["publishable"] and r["issues"]["UNRECOGNIZED_ATTRIBUTE"] == 1


def test_introduction_version_and_test_type_never_imply_support():
    b, _, _, _ = parse([application(**{"VERS.": "V16.0", "TIPO DE TESTE": "MODO COLABORATIVO"})])
    assert b.records[0].application_introduced_version == "V16.0"
    assert b.records[0].test_type == "MODO COLABORATIVO"
    assert b.records[0].release == ""
    assert b.motorcycles[0].status == "SEM_STATUS"


def test_primary_formula_rejected_even_if_optional():
    _, _, r, _ = parse([application(), application(**{"VÍDEO": "=1+1"})])
    assert not r["publishable"] and r["invalid_records"] == 1


def test_application_diff_cable_attributes_and_multiple_options():
    a = parse([application()])[0].records
    b = parse([application(CABO="MX-C2")])[0].records
    stats, details = application_differences(a, b)
    assert stats["applications"]["cable_changed"] == 1
    assert details[0][0] == "APPLICATION_CABLE_CHANGED"
    c = parse([application(CABO="MX-C2"), application(CABO="MX-C3")])[0].records
    stats, _ = application_differences(a, c)
    assert stats["applications"]["added"] == 2 and stats["applications"]["removed"] == 1
    d = parse([application(**{"VÍDEO": "NÃO"})])[0].records
    assert application_differences(a, d)[0]["applications"]["attributes_changed"] == 1
    assert application_differences(a, a)[0]["applications"]["unchanged"] == 1


@pytest.fixture
def scanner(tmp_path):
    path = tmp_path / "test.sqlite3"
    import_base(path, [motorcycle(status="SEM_SUPORTE"), motorcycle("F 900 R GT")])
    coverage(path, [ad(), ad(external_id="2", model="NEW 1234")], tmp_path / "coverage")
    return ScannerService(path)


def test_staging_hash_and_explicit_publication(scanner):
    payload = content([application(), application(SISTEMA="INJEÇÃO")])
    v = scanner.prepare(payload, "new.xlsx", "Fixture")["id"]
    with sqlite3.connect(scanner.database) as db:
        assert db.execute("select max(id) from imports").fetchone()[0] == 1
        assert db.execute("select content from scanner_versions where id=?", (v,)).fetchone()[0] == payload
    assert scanner.prepare(payload, "renamed.xlsx", "Fixture")["duplicate"]
    d = scanner.detail(v)
    with pytest.raises(ValueError):
        scanner.publish(v, "Fixture", "Conferido", False, d["parent_import_id"], d["impact_token"])
    result = publish(scanner, v)
    assert result["import_id"] == 2 and result["derived"] == "DONE"
    with sqlite3.connect(scanner.database) as db:
        assert db.execute("select count(*) from motorcycle_snapshots where import_id=2").fetchone()[0] == 1
        assert db.execute("select count(*) from system_records where import_id=2").fetchone()[0] == 2
        assert db.execute("select count(*) from motorcycle_snapshots where import_id=1").fetchone()[0] == 2
        assert db.execute("pragma integrity_check").fetchone()[0] == "ok"
        assert not db.execute("pragma foreign_key_check").fetchall()


def test_presence_preserved_absence_versioned_matching_per_vehicle(scanner):
    with ReviewRepository(scanner.database) as repo:
        items = repo.list_items()
        one = next(r for r in items if r["external_id"] == "1")
        two = next(r for r in items if r["external_id"] == "2")
        repo.decide(one["id"], "CONFIRMAR_MATCH", "Fixture", "Identidade conferida", "BMW|F900R|2025")
        repo.decide(two["id"], "NAO_EXISTE_NA_BASE", "Fixture", "Ausência conferida", None)
    v = scanner.prepare(
        content([application(), application(SISTEMA="INJEÇÃO"), application(MODELO="NEW 1234")]), "new.xlsx", "Fixture"
    )["id"]
    publish(scanner, v)
    with ReviewRepository(scanner.database) as repo:
        first, second = repo.show(one["id"]), repo.show(two["id"])
        assert first["effective"]["match_type"] == "EXATO_CONFIRMADO_HUMANAMENTE"
        assert first["effective"]["scanner_status"] == "SEM_STATUS"
        assert second["memory"]["stale_reason"] == "STALE_BASE_VERSION"
        assert len(second["history"]) == 1
    service = DashboardService(DashboardConfig(scanner.database))
    assert len(service.scanner("MX-C1")) == 2
    assert len(service.scanner("INJEÇÃO")) == 1
    assert len(service.scanner_applications("BMW|F900R|2025")) == 2
    now = datetime.now(timezone.utc) + timedelta(seconds=1)
    stats = ManagementMetricsService(scanner.database, clock=lambda: now).report(
        now.astimezone(ZONE).date(), now.astimezone(ZONE).date()
    )["base_stats"]
    assert stats["Total"] == 2 and stats["Aplicações (linhas)"] == 3 and stats["Sistemas distintos"] == 2


def test_publication_retains_unaffected_priority_and_manual_override(tmp_path, monkeypatch):
    from test_prioritization import priority as fixture

    from database.prioritization_repository import operational_sources

    p = fixture.__wrapped__(tmp_path)
    p.refresh()
    row = p.listing()["items"][0]
    p.override(
        row["id"], "low", author="Fixture", reason="Conferido", request_key="v16-manual", revision=row["revision"]
    )
    sources = operational_sources(p.config.database, "wr_motos")
    for source in sources:
        source["scanner_base_version"] += 1
    monkeypatch.setattr("services.prioritization_service.operational_sources", lambda *args: deepcopy(sources))
    result = p.refresh("scanner_publication")
    assert result["evaluated"] == 0 and result["retained"] == 10
    assert p.detail(row["id"])["effective_priority"] == "low"
    sources[0]["coverage"] = "SUPORTE_PARCIAL"
    result = p.refresh("scanner_publication")
    assert result["evaluated"] == 1 and result["retained"] == 9


def test_development_not_completed_when_application_added(scanner):
    from test_development import create

    from services.development_service import DevelopmentService

    dev = DevelopmentService(DashboardConfig(scanner.database))
    item = create(dev)
    v = scanner.prepare(content(), "new.xlsx", "Fixture")["id"]
    publish(scanner, v)
    assert dev.detail(item)["status"] == "NEW"


def test_new_format_transaction_rollback(scanner, monkeypatch):
    v = scanner.prepare(content(), "new.xlsx", "Fixture")["id"]

    def fail(*args):
        raise RuntimeError("failure after import")

    monkeypatch.setattr(scanner, "_alerts", fail)
    with pytest.raises(RuntimeError):
        publish(scanner, v)
    with sqlite3.connect(scanner.database) as db:
        assert db.execute("SELECT MAX(id) FROM imports").fetchone()[0] == 1
        assert db.execute("SELECT COUNT(*) FROM system_records WHERE import_id=2").fetchone()[0] == 0
    assert scanner.detail(v)["status"] == "VALIDATED"


def test_published_hash_cannot_be_published_twice(scanner):
    payload = content()
    v = scanner.prepare(payload, "new.xlsx", "Fixture")["id"]
    publish(scanner, v)
    assert scanner.prepare(payload, "copy.xlsx", "Fixture") == {"id": v, "duplicate": True}
    with pytest.raises(ValueError):
        publish(scanner, v)


def test_raw_identity_and_application_versions_preserved():
    base, _, _, _ = parse([application(MONTADORA=" bmw ", MODELO="F-900 R", **{"VERS.": "V3.0"})])
    assert base.records[0].raw_fields["manufacturer"] == " bmw "
    assert base.records[0].raw_fields["model"] == "F-900 R"
    assert base.records[0].key == "BMW|F900R|2025"
    assert base.records[0].application_introduced_version == "V3.0"


def test_multiple_thousand_applications_still_one_vehicle():
    base, _, report, _ = parse([application(SISTEMA=f"SISTEMA {i}") for i in range(1200)])
    assert len(base.motorcycles) == 1
    assert report["applications"] == 1200 and report["unique_systems"] == 1200


def test_new_format_dashboard_search_and_details(scanner, monkeypatch, all_navigation):
    from pathlib import Path

    from streamlit.testing.v1 import AppTest

    v = scanner.prepare(content([application(), application(SISTEMA="INJEÇÃO")]), "new.xlsx", "Fixture")["id"]
    monkeypatch.setenv("MOTO_DB", str(scanner.database))
    monkeypatch.setenv("MOTO_READ_ONLY", "1")
    ui = AppTest.from_file(str(Path("app/dashboard.py").resolve()), default_timeout=30).run()
    ui.radio(key="navigation").set_value("Atualização da base do scanner").run()
    ui.selectbox(key="scanner_version_choice").set_value(v).run()
    assert not ui.exception and not ui.error
    assert any("APLICAÇÃO GERAL" in str(i.value) for i in ui.info)
    publish(scanner, v)
    ui.radio(key="navigation").set_value("Base do scanner").run()
    next(t for t in ui.text_input if t.label == "Buscar na base").set_value("MX-C1").run()
    next(t for t in ui.selectbox if t.label == "Inspecionar identidade").set_value("BMW|F900R|2025").run()
    assert not ui.exception and not ui.error
    assert any("Aplicações do veículo" == t.value for t in ui.subheader)
    assert any("Cabo" in f.value.columns for f in ui.dataframe)
