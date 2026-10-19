"""Confirmed, atomic partner imports. Parsing and preview have no database writes."""

import csv
import hashlib
import io
import os
import re
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook

from database.inventory_repository import alert, event, now
from database.partner_repository import PartnerRepository
from database.repository import encode
from database.transaction import atomic_database
from partners.integration import integration
from partners.models import CollectionResult, PartnerMotorcycle
from partners.normalization import normalize_advertisement
from partners.sources import public_url
from services.coverage_service import build_coverage
from services.pipeline_lock import ExecutionLock

FIELDS = ("external_id", "manufacturer", "model", "version", "year", "price", "mileage", "detail_url", "image_url")
ALIASES = {
    "codigo": "external_id",
    "código": "external_id",
    "codigo de estoque": "external_id",
    "marca": "manufacturer",
    "modelo": "model",
    "ano": "year",
    "ano fabricação": "year",
    "versão": "version",
    "preço": "price",
    "quilometragem": "mileage",
    "url": "detail_url",
    "imagem": "image_url",
}


def writable(database, read_only=False):
    if read_only or os.environ.get("MOTO_READ_ONLY", "0").lower() in {"1", "true", "yes"}:
        raise ValueError("Modo somente leitura: operação desabilitada")
    if not Path(database).is_file():
        raise ValueError("Importe a base do scanner antes desta operação")


def read_file(content, filename):
    if not content or len(content) > 10_000_000:
        raise ValueError("Arquivo vazio ou maior que 10 MB")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    elif suffix == ".xlsx":
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 50_000_000:
                    raise ValueError("Planilha descompactada excede limite")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        except (zipfile.BadZipFile, KeyError, ParseError, OSError):
            raise ValueError("Planilha XLSX inválida ou corrompida") from None
        try:
            rows = []
            for cells in book.active.iter_rows():
                if len(rows) > 5000 or len(cells) > 100:
                    raise ValueError("Limite de linhas/colunas excedido")
                if any(c.data_type == "f" for c in cells):
                    raise ValueError("Use valores, sem fórmulas na planilha")
                rows.append([c.value for c in cells])
        except (zipfile.BadZipFile, KeyError, ParseError, OSError):
            raise ValueError("Planilha XLSX inválida ou corrompida") from None
        finally:
            book.close()
    else:
        raise ValueError("Formato permitido: CSV ou XLSX")
    if not rows or len(rows) > 5001 or len(rows[0]) > 100:
        raise ValueError("Arquivo deve ter cabeçalho e até 5.000 linhas")
    headers = [str(v or "").strip() for v in rows[0]]
    if not all(headers) or len(set(headers)) != len(headers):
        raise ValueError("Cabeçalhos vazios ou duplicados")
    data = []
    for values in rows[1:]:
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        if len(values) > len(headers):
            raise ValueError("Linha possui mais colunas que o cabeçalho")
        data.append(dict(zip(headers, values + [None] * (len(headers) - len(values)))))
    return headers, data


def default_mapping(headers):
    return {
        field: next((h for h in headers if h.lower() == field or ALIASES.get(h.lower()) == field), "")
        for field in FIELDS
    }


def clean_row(row):
    result = {k: str(row.get(k) if row.get(k) is not None else "").strip() for k in FIELDS}
    errors = []
    for k in ("manufacturer", "model"):
        if not result[k] or len(result[k]) > 200:
            errors.append("Sem marca" if k == "manufacturer" else "Sem modelo")
    year = result["year"]
    if year.endswith(".0"):
        year = year[:-2]
    if not re.fullmatch("[0-9]{4}", year) or not 1900 <= int(year) <= datetime.now().year + 2:
        errors.append("Sem ano válido")
    else:
        result["year"] = int(year)
    for key in ("detail_url", "image_url"):
        if result[key]:
            try:
                public_url(result[key])
            except ValueError:
                errors.append("URL inválida: " + key)
    if result["mileage"]:
        value = result["mileage"].removesuffix(".0")
        if not value.isascii() or not value.isdigit():
            errors.append("Quilometragem inválida")
        else:
            result["mileage"] = int(value)
    else:
        result["mileage"] = None
    if result["price"] and not re.fullmatch(r"[0-9]+(?:[.,][0-9]{1,2})?", result["price"]):
        errors.append("Preço inválido; use valor decimal sem separador de milhar")
    result["price"] = result["price"].replace(",", ".") or None
    if any(len(str(v or "")) > 2000 for v in result.values()):
        errors.append("Campo excede limite")
    return result, errors


def preview_file(partner, content, filename, mapping=None, mode="INCREMENTAL"):
    if not integration(partner).supports_manual_import:
        raise ValueError("Parceiro não aceita importação de estoque")
    if mode not in {"INCREMENTAL", "FULL_SNAPSHOT"}:
        raise ValueError("Modo de importação inválido")
    headers, raw = read_file(content, filename)
    mapping = default_mapping(headers) if mapping is None else dict(mapping)
    if any(k not in FIELDS or (v and v not in headers) for k, v in mapping.items()):
        raise ValueError("Mapeamento contém coluna desconhecida")
    selected = [v for v in mapping.values() if v]
    if len(set(selected)) != len(selected):
        raise ValueError("Uma coluna não pode mapear dois campos")
    digest = hashlib.sha256(content).hexdigest()
    rows = []
    seen = set()
    duplicates = missing_id = missing_year = missing_model = 0
    for index, data in enumerate(raw, 2):
        row, errors = clean_row({key: data.get(mapping.get(key, "")) for key in FIELDS})
        missing_model += int(not row["model"])
        missing_year += int(not isinstance(row["year"], int))
        synthetic = False
        if not row["external_id"]:
            missing_id += 1
            row["external_id"] = row["detail_url"] or f"file:{digest}:{index}"
            synthetic = not bool(row["detail_url"])
        if row["external_id"] in seen:
            duplicates += 1
            errors.append("ID duplicado no arquivo")
        seen.add(row["external_id"])
        if synthetic and mode == "FULL_SNAPSHOT":
            errors.append("Retrato completo exige código estável ou URL por anúncio")
        rows.append({"line": index, "data": row, "errors": errors, "generated_id": synthetic})
    valid = [r for r in rows if not r["errors"]]
    return {
        "partner": partner,
        "filename": Path(str(filename).replace("\\", "/")).name,
        "file_hash": digest,
        "mapping": mapping,
        "mode": mode,
        "total_rows": len(rows),
        "valid_rows": len(valid),
        "invalid_rows": len(rows) - len(valid),
        "duplicates": duplicates,
        "missing_id": missing_id,
        "missing_year": missing_year,
        "missing_model": missing_model,
        "rows": rows,
        "preview": rows[:20],
    }


def advertisement(partner, row, method):
    stamp = now()
    return normalize_advertisement(
        PartnerMotorcycle(
            partner,
            str(row["external_id"]),
            row["manufacturer"],
            row["model"],
            row.get("version") or None,
            row["year"],
            f"{row['manufacturer']} {row['model']}",
            row.get("detail_url") or "",
            stamp,
            "",
            None,
            price=row.get("price"),
            mileage=row.get("mileage"),
            raw_data={"integration_method": method},
            primary_image_url=row.get("image_url") or None,
        ),
        {},
    )


def publish_rows(database, partner, rows, complete, method, destination):
    result = CollectionResult(
        partner, now(), [advertisement(partner, r, method) for r in rows], complete=complete, method=method
    )
    with PartnerRepository(database) as repo:
        collection_id = repo.save_collection(result)
        if complete:
            repo.connection.execute(
                "UPDATE review_items SET active=0 WHERE partner=? AND external_id IN (SELECT external_id FROM partner_advertisements WHERE partner=? AND not_seen_in_latest_collection=1)",
                (partner, partner),
            )
    report = build_coverage(collection_id, database, Path(destination) / f"{partner}-{collection_id}")
    return collection_id, report


class PartnerInventoryImportService:
    def __init__(self, database, *, read_only=False, reports=None):
        self.database = Path(database)
        self.read_only = read_only
        self.reports = Path(reports or self.database.parent / "partner-import-reports")

    def preview(self, partner, content, filename, **kwargs):
        return preview_file(partner, content, filename, **kwargs)

    def confirm(
        self,
        partner,
        content,
        filename,
        *,
        reviewer,
        confirmed=False,
        authorized=False,
        motorcycles_confirmed=False,
        mapping=None,
        mode="INCREMENTAL",
    ):
        writable(self.database, self.read_only)
        if not confirmed or not authorized or not motorcycles_confirmed or not reviewer.strip():
            raise ValueError("Confirme importação, autorização de origem, motocicletas e responsável")
        try:
            review = preview_file(partner, content, filename, mapping, mode)
        except (ValueError, OSError):
            integration(partner)
            self.failure(partner, "IMPORT_FAILED", reviewer)
            raise
        if not review["valid_rows"] or (mode == "FULL_SNAPSHOT" and review["invalid_rows"]):
            self.failure(partner, "IMPORT_INVALID_ROWS", reviewer)
            raise ValueError("Sem linhas válidas ou retrato completo contém linhas inválidas")
        batch = hashlib.sha256(encode([review["file_hash"], review["mapping"], mode]).encode()).hexdigest()
        try:
            with (
                ExecutionLock(self.database, partner=partner),
                atomic_database(self.database),
                PartnerRepository(self.database) as repo,
            ):
                db = repo.connection
                previous = db.execute(
                    "SELECT id,collection_id FROM partner_inventory_imports WHERE partner=? AND batch_key=?",
                    (partner, batch),
                ).fetchone()
                if previous:
                    return {"id": previous[0], "collection_id": previous[1], "reused": True}
                if db.execute("SELECT MAX(id) FROM imports").fetchone()[0] is None:
                    raise ValueError("Importe a base do scanner primeiro")
                cid, coverage = publish_rows(
                    self.database,
                    partner,
                    [r["data"] for r in review["rows"] if not r["errors"]],
                    mode == "FULL_SNAPSHOT",
                    "FILE_IMPORT",
                    self.reports,
                )
                cur = db.execute(
                    "INSERT INTO partner_inventory_imports(partner,created_at,filename,file_hash,batch_key,mode,total_rows,valid_rows,invalid_rows,imported_rows,duplicates,user_action,reviewer,status,collection_id) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (
                        partner,
                        now(),
                        review["filename"],
                        review["file_hash"],
                        batch,
                        mode,
                        review["total_rows"],
                        review["valid_rows"],
                        review["invalid_rows"],
                        review["valid_rows"],
                        review["duplicates"],
                        "CONFIRM_IMPORT",
                        reviewer.strip(),
                        "IMPORTED",
                        cid,
                    ),
                )
                db.execute(
                    "INSERT INTO partner_import_mappings VALUES (?,?,?) ON CONFLICT(partner) DO UPDATE SET mapping_json=excluded.mapping_json,updated_at=excluded.updated_at",
                    (partner, encode(review["mapping"]), now()),
                )
                event(
                    db,
                    partner,
                    "IMPORT_CONFIRMED",
                    reviewer,
                    {
                        "import_id": cur.lastrowid,
                        "file_hash": review["file_hash"],
                        "mode": mode,
                        "valid_rows": review["valid_rows"],
                    },
                )
                alert(db, partner, "IMPORT_INVALID_ROWS", bool(review["invalid_rows"]))
                alert(db, partner, "IMPORT_PENDING", False)
                alert(db, partner, "IMPORT_FAILED", False)
                output = {
                    "id": cur.lastrowid,
                    "collection_id": cid,
                    "reused": False,
                    "coverage_counts": coverage["counts"],
                    "imported_rows": review["valid_rows"],
                }
            from services.prioritization_service import safe_refresh

            safe_refresh(self.database, partner, "collection")
            return output
        except Exception:
            self.failure(partner, "IMPORT_FAILED", reviewer)
            raise ValueError("Importação não concluída; estoque anterior preservado. Consulte a auditoria.") from None

    def failure(self, partner, code, reviewer):
        with PartnerRepository(self.database) as repo, repo.connection:
            alert(repo.connection, partner, code)
            event(repo.connection, partner, code, reviewer, {"status": "FAILED"})


def template_csv():
    return (",".join(FIELDS) + "\n").encode("utf-8-sig")
